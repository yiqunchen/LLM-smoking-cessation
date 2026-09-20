#!/usr/bin/env python3
"""Create the Word report for the prompt ablations (three domains, all metrics, bootstrap CIs, pairwise verdicts)."""
from __future__ import annotations

from pathlib import Path

import pandas as pd
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTDIR = ROOT / "figures" / "prompt_ablations"
OUTPUT = ROOT / "reports" / "prompt_ablation_results.docx"
MODEL_ORDER = ["GPT-4o-mini", "GPT-5", "DeepSeek-R1", "Grok-4.3", "Gemini-2.5-Pro"]
CONDITION_ORDER = [
    "PP + history + CBT/ACT",
    "PP + history (no CBT/ACT)",
    "History + ratings only",
    "History text only",
]
DOMAIN_ORDER = ["Content", "Coping", "Quitting"]
STANDARD_METRICS = ["accuracy", "macro_f1", "qwk"]
DIRECTIONAL_METRICS = ["directional_accuracy", "directional_macro_f1"]


def shade(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, text: str, bold: bool = False, color: str | None = None) -> None:
    cell.text = ""
    run = cell.paragraphs[0].add_run(text)
    run.bold = bold
    run.font.size = Pt(7.8)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def add_table(doc: Document, columns: list[str], rows: list[list[str]]) -> None:
    table = doc.add_table(rows=1, cols=len(columns))
    table.style = "Table Grid"
    for cell, title in zip(table.rows[0].cells, columns):
        set_cell_text(cell, title, bold=True, color="FFFFFF")
        shade(cell, "1F4E79")
    for row in rows:
        cells = table.add_row().cells
        for cell, value in zip(cells, row):
            set_cell_text(cell, value)
    doc.add_paragraph()


def format_ci(values: pd.Series, value: str) -> str:
    if value == "difference":
        return f"{values[value]:+.3f} [{values['ci_low']:+.3f}, {values['ci_high']:+.3f}]"
    return f"{values[value]:.3f} [{values['ci_low']:.3f}, {values['ci_high']:.3f}]"


def add_caption(doc: Document, text: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run(text)
    run.italic = True
    run.font.size = Pt(9)


def make_table_rows(frame: pd.DataFrame, index: list[str], metrics: list[str], value: str) -> list[list[str]]:
    pivot = frame.pivot(index=index, columns="metric", values=[value, "ci_low", "ci_high"])
    pivot = pivot.reindex(columns=pd.MultiIndex.from_product([[value, "ci_low", "ci_high"], metrics]))
    rows: list[list[str]] = []
    for keys, values in pivot.iterrows():
        cells = [str(key) for key in (keys if isinstance(keys, tuple) else (keys,))]
        for metric in metrics:
            row = pd.Series({
                value: values[(value, metric)],
                "ci_low": values[("ci_low", metric)],
                "ci_high": values[("ci_high", metric)],
            })
            cells.append(format_ci(row, value))
        rows.append(cells)
    return rows


PAIR_SHORT = {"PP + history + CBT/ACT": "Full PP", "PP + history (no CBT/ACT)": "No CBT/ACT",
              "History + ratings only": "Hist + ratings", "History text only": "Hist text"}
METRIC_LABEL = {"accuracy": "Exact accuracy", "macro_f1": "Macro-F1", "qwk": "QWK",
                "directional_accuracy": "Directional accuracy", "directional_macro_f1": "Directional macro-F1"}


def pairwise_table_rows(pairwise: pd.DataFrame, metric: str, completed: list[str]) -> tuple[list[str], list[list[str]]]:
    """Rows = model x domain; columns = the six condition pairs; cell = difference [95% CI] with a verdict marker."""
    block = pairwise[pairwise["metric"] == metric]
    pairs = list(dict.fromkeys(zip(block["reference"], block["comparison"])))
    columns = ["Model", "Domain"] + [f"{PAIR_SHORT[c]} − {PAIR_SHORT[r]}" for r, c in pairs]
    rows = []
    for model in completed:
        for domain in DOMAIN_ORDER:
            cells = [model, domain]
            for reference, comparison in pairs:
                hit = block[(block["model"] == model) & (block["domain"] == domain)
                            & (block["reference"] == reference) & (block["comparison"] == comparison)].iloc[0]
                marker = "" if hit["verdict"] == "comparable" else " *"
                cells.append(f"{hit['difference']:+.3f} [{hit['ci_low']:+.3f}, {hit['ci_high']:+.3f}]{marker}")
            rows.append(cells)
    return columns, rows


def main() -> None:
    metrics_path = OUTDIR / "prompt_ablation_bootstrap_metrics_dt10.csv"
    deltas_path = OUTDIR / "prompt_ablation_bootstrap_deltas_vs_pp_cbtact_dt10.csv"
    pairwise_path = OUTDIR / "prompt_ablation_bootstrap_pairwise_dt10.csv"
    prediction_path = OUTDIR / "prompt_ablation_prediction_summary_dt10.csv"
    audit_path = OUTDIR / "prompt_ablation_integrity_audit_dt10.csv"
    for path in (metrics_path, deltas_path, pairwise_path, prediction_path, audit_path):
        if not path.exists():
            raise SystemExit(f"Missing source table: {path.name}; run audit, bootstrap, and plot scripts first")
    metrics = pd.read_csv(metrics_path)
    deltas = pd.read_csv(deltas_path)
    pairwise = pd.read_csv(pairwise_path)
    prediction = pd.read_csv(prediction_path)
    audit = pd.read_csv(audit_path)
    completed = [model for model in MODEL_ORDER if model in set(metrics["model"])]
    for frame, condition_column in ((metrics, "condition"), (deltas, "comparison")):
        frame["model"] = pd.Categorical(frame["model"], categories=completed, ordered=True)
        frame["domain"] = pd.Categorical(frame["domain"], categories=DOMAIN_ORDER, ordered=True)
        frame[condition_column] = pd.Categorical(frame[condition_column], categories=CONDITION_ORDER, ordered=True)
    for frame in (prediction,):
        frame["model"] = pd.Categorical(frame["model"], categories=completed, ordered=True)
        frame["domain"] = pd.Categorical(frame["domain"], categories=DOMAIN_ORDER, ordered=True)

    document = Document()
    section = document.sections[0]
    section.orientation = WD_ORIENT.LANDSCAPE
    section.page_width, section.page_height = section.page_height, section.page_width
    for side in ("top_margin", "bottom_margin", "left_margin", "right_margin"):
        setattr(section, side, Inches(.55))
    normal = document.styles["Normal"]
    normal.font.name = "Helvetica"
    normal.font.size = Pt(9)

    title = document.add_heading("Prompt-Component Sensitivity Analysis", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph(
        "All results use the canonical dt10-k7 held-out test set (898 messages from 301 participants). "
        "The evaluated outcomes are Content, Coping, and Quitting, each reported separately; no cross-domain mean is calculated."
    )
    document.add_paragraph(
        "Intervals are 95% percentile intervals from 2,000 participant-clustered bootstrap replicates; paired differences use the same "
        "replicate for both configurations. Directional metrics use low (1–2), neutral (3), and high (4–5). "
        "No multiplicity correction is applied to the pairwise comparisons."
    )
    document.add_paragraph("Configurations: " + "; ".join(f"{k} = {v}" for k, v in PAIR_SHORT.items()) + ".")

    # ---- performance ------------------------------------------------------
    document.add_heading("Table 1. Exact and ordinal performance by domain", level=1)
    standard = metrics[metrics["metric"].isin(STANDARD_METRICS)].sort_values(["domain", "model", "condition"])
    add_table(document, ["Domain", "Model", "Configuration", "Exact accuracy (95% CI)", "Macro-F1 (95% CI)", "QWK (95% CI)"],
              make_table_rows(standard, ["domain", "model", "condition"], STANDARD_METRICS, "estimate"))
    document.add_picture(str(OUTDIR / "prompt_ablation_metrics_by_domain_dt10.png"), width=Inches(9.0))
    add_caption(document, "Figure 1. Exact accuracy, macro-F1, and QWK by rating domain. Lines connect configurations evaluated on the same canonical test messages; "
                          "95% intervals are given in Table 1.")

    document.add_heading("Table 2. Directional performance by domain", level=1)
    directional = metrics[metrics["metric"].isin(DIRECTIONAL_METRICS)].sort_values(["domain", "model", "condition"])
    add_table(document, ["Domain", "Model", "Configuration", "Directional accuracy (95% CI)", "Directional macro-F1 (95% CI)"],
              make_table_rows(directional, ["domain", "model", "condition"], DIRECTIONAL_METRICS, "estimate"))
    document.add_picture(str(OUTDIR / "prompt_ablation_directional_by_domain_dt10.png"), width=Inches(9.0))
    add_caption(document, "Figure 2. Directional accuracy and directional macro-F1 by rating domain; 95% intervals are given in Table 2.")

    # ---- differences from the full PP prompt ------------------------------
    document.add_heading("Table 3. Paired differences from the full PP prompt (exact and ordinal metrics)", level=1)
    standard_delta = deltas[deltas["metric"].isin(STANDARD_METRICS)].sort_values(["domain", "model", "comparison"])
    add_table(document, ["Domain", "Model", "Compared configuration", "Δ accuracy (95% CI)", "Δ macro-F1 (95% CI)", "Δ QWK (95% CI)"],
              make_table_rows(standard_delta, ["domain", "model", "comparison"], STANDARD_METRICS, "difference"))
    document.add_heading("Table 4. Paired differences from the full PP prompt (directional metrics)", level=1)
    directional_delta = deltas[deltas["metric"].isin(DIRECTIONAL_METRICS)].sort_values(["domain", "model", "comparison"])
    add_table(document, ["Domain", "Model", "Compared configuration", "Δ directional accuracy (95% CI)", "Δ directional macro-F1 (95% CI)"],
              make_table_rows(directional_delta, ["domain", "model", "comparison"], DIRECTIONAL_METRICS, "difference"))
    document.add_paragraph("Positive differences favor the compared configuration.")

    # ---- all pairwise comparisons, one table per metric -------------------
    table_number = 5
    for metric in STANDARD_METRICS + DIRECTIONAL_METRICS:
        document.add_heading(f"Table {table_number}. All pairwise configuration differences: {METRIC_LABEL[metric]}", level=1)
        columns, rows = pairwise_table_rows(pairwise, metric, completed)
        add_table(document, columns, rows)
        table_number += 1
    document.add_paragraph("Each cell is comparison minus reference with its paired 95% interval; * marks intervals that exclude zero.")
    document.add_picture(str(OUTDIR / "prompt_ablation_pairwise_qwk_dt10.png"), width=Inches(9.0))
    add_caption(document, "Figure 3. Paired QWK differences for all six configuration pairs. Shaded cells have 95% intervals excluding zero; white cells are statistically comparable.")

    # ---- distributions ----------------------------------------------------
    document.add_heading(f"Table {table_number}. Distribution of predicted ratings", level=1)
    document.add_paragraph(
        "Mean and SD of observed and predicted ratings, bias (predicted minus observed), mean absolute error, and the share of "
        "messages predicted exactly or within one level, per configuration and domain."
    )
    prediction_rows = []
    for row in prediction.sort_values(["domain", "model", "condition_key"]).itertuples():
        prediction_rows.append([str(row.domain), str(row.model), str(row.condition), f"{row.observed_mean:.2f} ({row.observed_sd:.2f})",
                                f"{row.predicted_mean:.2f} ({row.predicted_sd:.2f})", f"{row.bias_predicted_minus_observed:+.2f}",
                                f"{row.mean_absolute_error:.2f}", f"{row.exact_rate:.3f}", f"{row.within_one_rate:.3f}"])
    add_table(document, ["Domain", "Model", "Configuration", "Observed mean (SD)", "Predicted mean (SD)", "Bias", "MAE", "Exact", "Within 1"], prediction_rows)
    table_number += 1
    document.add_picture(str(OUTDIR / "prompt_ablation_rating_distributions_dt10.png"), width=Inches(8.4))
    add_caption(document, "Figure 4. Human versus predicted rating distributions on the 898 held-out messages, by prompt configuration and domain. "
                          "Gray background bars are the human ratings; coloured bars are each model's predictions.")

    # ---- data integrity ---------------------------------------------------
    document.add_heading(f"Table {table_number}. Result-file integrity audit", level=1)
    document.add_paragraph(
        "Every saved row was checked against the canonical test set for its index, participant, message text, ground truth, "
        "allowed prediction label, condition identity, and seven-message history count."
    )
    audit_rows = [[str(r.model), str(r.condition_key), f"{r.n_saved_rows}/{r.n_expected_rows}", str(r.status), str(r.file_sha256)[:16] + "…"]
                  for r in audit.itertuples()]
    add_table(document, ["Model", "Configuration", "Rows", "Status", "File SHA-256 (prefix)"], audit_rows)

    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    document.save(OUTPUT)
    print(f"Wrote {OUTPUT}")


if __name__ == "__main__":
    main()
